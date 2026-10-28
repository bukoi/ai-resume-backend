from __future__ import annotations

import calendar
import re
from dataclasses import dataclass, replace
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from app.schema.resume import Resume
from app.schema.template import RoleStyle, SectionConfig, Template

PAGE_SIZES = {"A4": (8.27, 11.69), "Letter": (8.5, 11.0), "Legal": (8.5, 14.0)}
DEFAULT_SIZES = {"name": 20, "contact": 9, "section_title": 11, "entry_title": 10,
                 "entry_subtitle": 9, "body": 10, "bullet": 10, "date": 9, "link": 10}
DEFAULT_HEADINGS = {"summary": "Professional Summary", "skills": "Skills", "experience": "Experience",
                    "projects": "Projects", "certifications": "Certifications",
                    "achievements": "Achievements", "education": "Education", "languages": "Languages"}
FONT_ALIASES = {"arialmt": "Arial", "arialboldmt": "Arial", "timesnewromanpsmt": "Times New Roman"}
SEPARATORS = {"|": " | ", ",": ", ", "–": " – ", "-": " - ", "•": " • ", "·": " · ",
              "at": " at ", "to": " to ", ":": ": "}
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
MONTHS |= {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}
MONTHS["sept"] = 9
PRESENT = {"present", "current", "now", "ongoing", "till date", "to date"}


# ---------- small text helpers ----------

def tidy(s: str | None) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def hexcolor(c: str | None) -> str | None:
    return c.lstrip("#").upper() if c else None


def norm_font(name: str, fallback: str) -> str:
    key = re.sub(r"[\s\-_]", "", name or "").lower()
    return FONT_ALIASES.get(key, name or fallback)


def clean_url(url: str) -> str:
    parts = urlsplit(url.strip())
    query = [(k, v) for k, v in parse_qsl(parts.query) if not k.lower().startswith("utm_")]
    return urlunsplit(parts._replace(query=urlencode(query)))


def short_url(url: str) -> str:
    parts = urlsplit(clean_url(url))
    text = parts.netloc.removeprefix("www.") + parts.path
    return text.rstrip("/") + (f"?{parts.query}" if parts.query else "")


def location_text(loc: str | None, detail: str) -> str:
    loc = re.sub(r",(?=\S)", ", ", tidy(loc))
    if not loc or detail == "hidden":
        return ""
    parts = [p.strip() for p in loc.split(",")]
    n = {"city": 1, "city_region": 2}.get(detail)
    return ", ".join(parts[:n]) if n else loc


def fmt_date(s: str | None, t: Template) -> str:
    s = tidy(s)
    if not s:
        return ""
    if s.lower() in PRESENT:
        return t.dates.present_label
    fmt = t.dates.format
    if fmt is None:
        return s                                   # keep dates as written
    month = year = None
    m = re.fullmatch(r"([A-Za-z]+)\.?\s+(\d{4})", s)
    if m and m.group(1).lower() in MONTHS:
        month, year = MONTHS[m.group(1).lower()], int(m.group(2))
    else:
        m = re.fullmatch(r"(\d{1,2})[./](\d{4})", s)
        if m and 1 <= int(m.group(1)) <= 12:
            month, year = int(m.group(1)), int(m.group(2))
    if year is None:
        return s                                   # year-only or unknown: leave alone
    return {"long_name": f"{calendar.month_name[month]} {year}",
            "short_name": f"{calendar.month_abbr[month]} {year}",
            "MM/YYYY": f"{month:02d}/{year}", "MM.YYYY": f"{month:02d}.{year}",
            "YYYY": str(year)}[fmt]


def date_range(start: str | None, end: str | None, t: Template) -> str:
    parts = [fmt_date(x, t) for x in (start, end) if x]
    return SEPARATORS[t.dates.range_separator].join(p for p in parts if p)


def add_bottom_rule(p, color: str, width_pt: float, style: str = "single", space: float | None = None) -> None:
    """Must run before other paragraph properties so the XML order stays valid."""
    pPr = p._p.get_or_add_pPr()
    pBdr, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    bottom.set(qn("w:val"), style or "single")
    bottom.set(qn("w:sz"), str(max(2, int(width_pt * 8))))
    bottom.set(qn("w:space"), str(int(space) if space is not None else 1))
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


@dataclass
class St:                       # a fully resolved text style
    font: str
    size: float
    bold: bool
    italic: bool
    caps: bool
    color: str | None
    align: str | None


@dataclass
class Seg:                      # one piece of an entry's title line
    text: str
    url: str | None = None
    sep: str = " | "            # separator placed before this piece
    kind: str = "title"         # title | link | meta


@dataclass
class Item:                     # one contact detail
    kind: str
    text: str
    url: str | None = None


class Renderer:
    def __init__(self, resume: Resume, template: Template):
        self.r, self.t = resume, template          # self.r.header = content, self.t.header = layout
        self.doc = Document()
        self.line_height = template.spacing.line_height

        pg = template.page
        w, h = (pg.width_in, pg.height_in) if pg.width_in and pg.height_in else PAGE_SIZES.get(pg.format, PAGE_SIZES["A4"])
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(w), Inches(h)
        margin = lambda v: 0.6 if v is None else v
        left, right = margin(pg.margin_left_in), margin(pg.margin_right_in)
        sec.left_margin, sec.right_margin = Inches(left), Inches(right)
        sec.top_margin, sec.bottom_margin = Inches(margin(pg.margin_top_in)), Inches(margin(pg.margin_bottom_in))
        sec.header_distance = Inches(min(pg.header_footer_in or 0.3, margin(pg.margin_top_in)))
        sec.footer_distance = Inches(min(pg.header_footer_in or 0.3, margin(pg.margin_bottom_in)))
        self.usable_in = w - left - right

        body = self.st("body")
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = body.font, Pt(body.size)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), body.font)
        self.num_id = self._bullet_numbering()

    # ---------- style, numbering, paragraph and run helpers ----------

    def st(self, name: str, base: str | None = None) -> St:
        ty = self.t.typography
        r: RoleStyle = getattr(ty, name)
        if base:
            s = replace(self.st(base))
        else:
            heading_like = name in ("name", "section_title", "entry_title")
            s = St(font=(ty.primary_font if heading_like else ty.secondary_font) or ty.fallback_font,
                   size=DEFAULT_SIZES[name], bold=False, italic=False, caps=False,
                   color=self.t.colors.text, align=None)
        if r.font: s.font = r.font
        if r.size_pt: s.size = r.size_pt
        if r.bold is not None: s.bold = r.bold
        if r.italic is not None: s.italic = r.italic
        if r.caps is not None: s.caps = r.caps
        if r.color: s.color = r.color
        if r.align: s.align = r.align
        s.font, s.color = norm_font(s.font, ty.fallback_font), hexcolor(s.color)
        return s

    def gap(self, name: str, default: float) -> float:
        v = getattr(self.t.spacing, name)
        return default if v is None else v

    def _bullet_numbering(self) -> int:
        b, ns = self.t.bullets, nsdecls("w")
        indent = 0.1 if b.indent_in is None else b.indent_in
        hang = 0.15 if b.hanging_in is None else b.hanging_in
        font = self.st("bullet", base="body").font
        abstract = parse_xml(
            f'<w:abstractNum {ns} w:abstractNumId="900">'
            '<w:multiLevelType w:val="hybridMultilevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            f'<w:lvlText w:val="{b.char}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{int((indent + hang) * 1440)}" w:hanging="{int(hang * 1440)}"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}" w:hint="default"/></w:rPr>'
            '</w:lvl></w:abstractNum>')
        numbering = self.doc.part.numbering_part.element
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)        # abstractNum elements must come before num
        else:
            numbering.append(abstract)
        return numbering.add_num(900).numId

    def para(self, *, before=0.0, after=0.0, align=None, keep_next=False, rule=None):
        p = self.doc.add_paragraph()
        if rule:
            add_bottom_rule(p, *rule)
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        if self.line_height:
            pf.line_spacing = self.line_height
        if align:
            p.alignment = ALIGN[align]
        if keep_next:
            pf.keep_with_next = True
        return p

    def run(self, p, text, st: St, *, bold=None, italic=None, color=None):
        r = p.add_run(text)
        f = r.font
        f.name, f.size = st.font, Pt(st.size)
        r._r.rPr.rFonts.set(qn("w:eastAsia"), st.font)
        r._r.rPr.rFonts.set(qn("w:cs"), st.font)
        f.bold = st.bold if bold is None else bold
        f.italic = st.italic if italic is None else italic
        if st.caps:
            f.all_caps = True
        c = hexcolor(color) or st.color
        if c:
            f.color.rgb = RGBColor.from_string(c)
        return r

    def link(self, p, text, url, st: St, *, bold=None, underline=False):
        r_id = p.part.relate_to(clean_url(url), RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        h.set(qn("w:history"), "1")
        p._p.append(h)
        r = self.run(p, text, st, bold=bold, color=self.t.colors.link)
        r.font.underline = underline
        h.append(r._r)                              # move the run inside the hyperlink

    def link_label(self, label: str, url: str | None) -> str:
        d = self.t.links.display
        if url and (d == "url" or not label):
            return short_url(url)
        if url and d == "label_and_url":
            return f"{label} ({short_url(url)})"
        return label

    def right_tab(self, p):
        p.paragraph_format.tab_stops.add_tab_stop(Inches(self.usable_in), WD_TAB_ALIGNMENT.RIGHT)

    def line(self, role: str, text: str, *, after=None):
        text = tidy(text)
        if text:
            p = self.para(after=self.gap("inside_block_pt", 1) if after is None else after)
            self.run(p, text, self.st(role))

    def item_para(self, as_bullet: bool):
        p = self.para(after=self.gap("inside_block_pt", 1))
        if as_bullet:
            numPr = p._p.get_or_add_pPr().get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = self.num_id
        return p

    def bullets(self, items):
        st = self.st("bullet", base="body")
        for text in items:
            if tidy(text):
                self.run(self.item_para(True), tidy(text), st)

    def links_line(self, links):
        p = self.para(after=self.gap("inside_block_pt", 1))
        st = self.st("link", base="entry_subtitle")
        for i, l in enumerate(links):
            if i:
                self.run(p, " | ", st)
            text = self.link_label(tidy(l.label), l.url)
            if l.url:
                self.link(p, text, l.url, st, underline=self.t.links.underline)
            else:
                self.run(p, text, st)

    def entry_head(self, segs: list[Seg], date_text: str, *, first: bool):
        title, meta = self.st("entry_title"), self.st("entry_subtitle")
        link, date_st = self.st("link", base="entry_title"), self.st("date", base="entry_subtitle")
        p = self.para(before=0 if first else self.gap("between_blocks_pt", 4), keep_next=True)
        for i, s in enumerate(segs):
            if i:
                self.run(p, s.sep, title)
            if s.kind == "link" and s.url:
                self.link(p, s.text, s.url, link, bold=False, underline=self.t.links.underline)
            else:
                self.run(p, s.text, meta if s.kind == "meta" else title)
        pos = self.t.dates.position
        if date_text and pos == "right_same_line":
            self.right_tab(p)
            self.run(p, "\t" + date_text, date_st)
        elif date_text and pos == "inline_same_line":
            self.run(p, " | " + date_text, date_st)
        elif date_text:
            self.run(self.para(keep_next=True), date_text, date_st)

    # ---------- header ----------

    def contact_items(self) -> list[Item]:
        h, hdr = self.r.header, self.t.header
        email = re.sub(r"\s+", "", h.email or "")
        buckets: dict[str, list[Item]] = {"location": [], "email": [], "phone": [], "links": []}
        loc = location_text(h.location, hdr.location_detail)
        if loc:
            buckets["location"].append(Item("location", loc))
        if email:
            buckets["email"].append(Item("email", email, f"mailto:{email}"))
        if tidy(h.phone):
            buckets["phone"].append(Item("phone", tidy(h.phone)))
        for l in h.links:
            url = (l.url or "").strip()
            if url.lower().startswith("mailto:") or re.sub(r"\s+", "", l.label or "").lower() == email.lower():
                continue                            # the email is already shown from header.email
            text = self.link_label(tidy(l.label), url or None)
            if text:
                buckets["links"].append(Item("link", text, url or None))
        return [it for key in hdr.contact_order for it in buckets[key]]

    def header(self):
        hdr, contact = self.t.header, self.st("contact")
        p_name = self.para(after=2, align=hdr.alignment)
        self.run(p_name, tidy(self.r.header.name), self.st("name"))
        items = self.contact_items()
        if hdr.contact_layout == "stacked":
            groups = [[i] for i in items]
        elif hdr.contact_layout == "two_lines":
            groups = [g for g in ([i for i in items if i.kind != "link"], [i for i in items if i.kind == "link"]) if g]
        else:
            groups = [items] if items else []
        last_p = p_name
        for group in groups:
            p = self.para(after=1, align=hdr.alignment)
            last_p = p
            for i, it in enumerate(group):
                if i:
                    self.run(p, SEPARATORS[hdr.contact_separator], contact)
                if it.url:
                    self.link(p, it.text, it.url, contact, underline=self.t.links.underline and it.kind == "link")
                else:
                    self.run(p, it.text, contact)

        if hdr.divider and hdr.divider.enabled and last_p is not None:
            div = hdr.divider
            rule_color = hexcolor(div.color or self.t.colors.accent) or "000000"
            add_bottom_rule(last_p, rule_color, div.width_pt or 0.75, style=div.style, space=div.space_pt)

    def heading(self, text: str):
        st, rule = self.st("section_title"), None
        if self.t.section_rule.enabled:
            sr = self.t.section_rule
            rule = (hexcolor(sr.color or self.t.colors.accent) or st.color or "000000",
                    sr.width_pt or 0.75, sr.style, sr.space_pt)
        p = self.para(before=self.gap("between_sections_pt", 8), after=self.gap("title_to_content_pt", 3),
                      align=st.align, keep_next=True, rule=rule)
        self.run(p, tidy(text), st)

    # ---------- sections (method names match the Resume keys) ----------

    def summary(self):
        self.run(self.para(), tidy(self.r.summary), self.st("body"))

    def skills(self):
        lay, st = self.t.section_layouts.skills, self.st("body")
        if lay.mode == "category_lines":
            for g in self.r.skills:
                p = self.para(after=self.gap("inside_block_pt", 1))
                if tidy(g.category):
                    sep = lay.category_separator
                    label = f"{tidy(g.category)}{sep} " if sep == ":" else f"{tidy(g.category)} {sep} "
                    self.run(p, label, st, bold=lay.category_bold)
                self.run(p, lay.item_separator.join(tidy(i) for i in g.items), st)
        else:        # "inline"; "columns" would need a table, which hurts ATS parsing
            flat = [tidy(i) for g in self.r.skills for i in g.items]
            self.run(self.para(), lay.item_separator.join(flat), st)

    def experience(self):
        lay = self.t.section_layouts.experience
        sep = SEPARATORS[lay.title_separator]
        for i, e in enumerate(self.r.experience):
            parts = [e.role, e.company] if lay.title_order == "role_first" else [e.company, e.role]
            segs = [Seg(tidy(x), sep=sep) for x in parts if tidy(x)]
            loc = tidy(e.location)
            if loc and lay.location_position == "inline_same_line":
                segs.append(Seg(loc, sep=sep))
            self.entry_head(segs, date_range(e.start, e.end, self.t), first=i == 0)
            if loc and lay.location_position in ("own_line", "right_same_line"):
                self.line("entry_subtitle", loc)
            self.bullets(e.bullets)

    def projects(self):
        lay = self.t.section_layouts.projects
        for i, pr in enumerate(self.r.projects):
            segs = [Seg(tidy(pr.name))]
            if tidy(pr.subtitle) and lay.subtitle_position == "inline":
                segs.append(Seg(tidy(pr.subtitle), sep=" – ", kind="meta"))
            tech = (f"{lay.tech_label} " if lay.tech_label else "") + lay.tech_separator.join(tidy(x) for x in pr.tech)
            if pr.tech and lay.tech_position == "inline":
                segs.append(Seg(tech, sep=" | ", kind="meta"))
            if lay.link_position in ("after_title", "title_line_right"):
                for l in pr.links:
                    segs.append(Seg(self.link_label(tidy(l.label), l.url), l.url, " | ", "link"))
            self.entry_head(segs, date_range(pr.start, pr.end, self.t), first=i == 0)
            if tidy(pr.subtitle) and lay.subtitle_position == "own_line":
                self.line("entry_subtitle", pr.subtitle)
            if pr.tech and lay.tech_position == "own_line":
                self.line("entry_subtitle", tech)
            if pr.links and lay.link_position == "own_line":
                self.links_line(pr.links)
            self.bullets(pr.bullets)
            if pr.tech and lay.tech_position == "after_bullets":
                self.line("entry_subtitle", tech)

    def education(self):
        lay = self.t.section_layouts.education
        for i, ed in enumerate(self.r.education):
            segs = [Seg(tidy(ed.degree))]
            if tidy(ed.institution) and lay.institution_position == "inline":
                segs.append(Seg(tidy(ed.institution), sep=", "))
            self.entry_head(segs, date_range(ed.start, ed.end, self.t), first=i == 0)
            bits = [tidy(ed.institution)] if tidy(ed.institution) and lay.institution_position == "own_line" else []
            if lay.show_location and tidy(ed.location):
                bits.append(tidy(ed.location))
            if lay.show_grade and tidy(ed.grade):
                bits.append(tidy(ed.grade))
            self.line("entry_subtitle", " | ".join(bits))
            if lay.notes_as_bullets:
                self.bullets(ed.notes)
            else:
                for n in ed.notes:
                    self.line("entry_subtitle", n)

    def certifications(self):
        lay, st = self.t.section_layouts.certifications, self.st("body")
        sep = SEPARATORS.get(lay.issuer_separator, f" {lay.issuer_separator} ")
        for c in self.r.certifications:
            p = self.item_para(lay.as_bullets)
            self.run(p, tidy(c.title), st, bold=True)
            rest = f"{sep}{tidy(c.issuer)}" if tidy(c.issuer) else ""
            if c.date and lay.date_position == "inline":
                rest += f" ({fmt_date(c.date, self.t)})"
            if tidy(c.details):
                rest += f" | {tidy(c.details)}"
            self.run(p, rest, st)
            if c.date and lay.date_position == "right_same_line":
                self.right_tab(p)
                self.run(p, "\t" + fmt_date(c.date, self.t), self.st("date", base="entry_subtitle"))

    def achievements(self):
        lay, st = self.t.section_layouts.achievements, self.st("body")
        sep = SEPARATORS.get(lay.detail_separator, f" {lay.detail_separator} ")
        for a in self.r.achievements:
            p = self.item_para(lay.as_bullets)
            self.run(p, tidy(a.title), st, bold=lay.bold_title)
            if tidy(a.detail):
                self.run(p, f"{sep}{tidy(a.detail)}", st)

    def languages(self):
        lay, st = self.t.section_layouts.languages, self.st("body")

        def one(l):
            name, level = tidy(l.language), tidy(l.proficiency)
            if not level:
                return name
            return {"parentheses": f"{name} ({level})", "dash": f"{name} – {level}",
                    "colon": f"{name}: {level}"}[lay.proficiency_style]

        items = [one(l) for l in self.r.languages]
        if lay.mode == "inline":
            self.run(self.para(), ", ".join(items), st)
        else:
            for text in items:
                self.run(self.item_para(True), text, st)

    def custom(self, cs):
        mode, lines = self.t.section_layouts.custom.mode, [tidy(x) for x in cs.lines if tidy(x)]
        st = self.st("body")
        if mode == "paragraph":
            self.run(self.para(), " ".join(lines), st)
        elif mode == "inline_list":
            self.run(self.para(), ", ".join(lines), st)
        else:
            self.bullets(lines)
        if cs.links:
            self.links_line(cs.links)

    # ---------- assembling ----------

    def has(self, key: str, custom_title: str | None = None) -> bool:
        r = self.r
        if key == "summary":
            return bool(tidy(r.summary))
        if key == "custom":
            return any(c.title == custom_title and (c.lines or c.links) for c in r.custom_sections)
        return bool(getattr(r, key, None))

    def sections(self) -> list[SectionConfig]:
        """The template's sections, plus any content the template forgot (so nothing vanishes)."""
        secs = list(self.t.sections)
        covered = {k for s in secs for k in (s.key, *s.also)}
        covered_custom = {s.custom_title for s in secs if s.key == "custom"}
        for key, heading in DEFAULT_HEADINGS.items():
            if key not in covered and self.has(key):
                secs.append(SectionConfig(key=key, heading=heading))
        for cs in self.r.custom_sections:
            if cs.title not in covered_custom:
                secs.append(SectionConfig(key="custom", heading=cs.title, custom_title=cs.title))
        return secs

    def build(self) -> Document:
        self.header()
        for s in self.sections():
            keys = [k for k in (s.key, *s.also) if s.visible and self.has(k, s.custom_title)]
            if not keys:
                continue                            # no empty headings
            self.heading(s.heading or DEFAULT_HEADINGS.get(s.key, ""))
            for k in keys:
                if k == "custom":
                    for cs in self.r.custom_sections:
                        if cs.title == s.custom_title:
                            self.custom(cs)
                else:
                    getattr(self, k)()
        props = self.doc.core_properties
        props.title, props.author, props.comments = f"{tidy(self.r.header.name)} - Resume", tidy(self.r.header.name), ""
        return self.doc


def render_docx(resume: Resume, template: Template, out_path: str | Path) -> Path:
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    Renderer(resume, template).build().save(out)
    return out
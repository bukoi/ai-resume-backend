from __future__ import annotations

import sys
from dataclasses import dataclass, field
from statistics import median

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
BULLET_CHARS = "•●▪■◦○▸►➢"


@dataclass
class Block:
    text: str
    where: str = "body"          # body | header | footer | textbox | table rXcY
    style: str = ""
    bullet: bool = False
    bold: bool = False           # all text bold
    italic: bool = False         # all text italic
    caps: bool = False           # ALL CAPS text
    size: float | None = None    # largest font size in pt
    align: str = ""
    rule: bool = False           # has a bottom border line
    links: list[tuple[str, str]] = field(default_factory=list)
    para: Paragraph | None = field(default=None, repr=False)


# ---------- formatting helpers (resolve style inheritance) ----------

def _style_attr(style, name):
    while style is not None:
        value = getattr(style.font, name)
        if value is not None:
            return value
        style = style.base_style
    return None


def _run_attr(run, p, name):
    value = getattr(run.font, name)
    return value if value is not None else _style_attr(p.style, name)


def _size(run, p):
    if run.font.size:
        return run.font.size.pt
    s = _style_attr(p.style, "size")
    return s.pt if s else None


def _has_bottom_rule(p) -> bool:
    candidates = [p._p.pPr]
    if p.style is not None:
        candidates.append(p.style.element.pPr)
    for pPr in candidates:
        if pPr is None:
            continue
        bdr = pPr.find(qn("w:pBdr"))
        if bdr is not None and bdr.find(qn("w:bottom")) is not None:
            return True
    return False

def _paragraph_text(p: Paragraph) -> str:
    return "".join(
        node.text or ""
        for node in p._p.iter(qn("w:t"))
    )

def _is_bullet(p) -> bool:
    pPr = p._p.pPr
    if pPr is not None and pPr.numPr is not None:
        return True
    style = p.style
    while style is not None:
        sp = style.element.pPr
        if sp is not None and sp.numPr is not None:
            return True
        style = style.base_style
    first = _paragraph_text(p).lstrip()[:1]
    return bool(first) and first in BULLET_CHARS


def _alignment(p) -> str:
    a = p.alignment
    if a is None and p.style is not None:
        a = p.style.paragraph_format.alignment
    return a.name.lower() if a is not None else ""


# ---------- walking the document in true reading order ----------

def _iter_items(container_el, parent):
    for child in container_el.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)
        elif child.tag == qn("w:sdt"):                 # content controls
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_items(content, parent)


def _add_paragraph(p: Paragraph, where: str, out: list[Block]) -> None:
    text = _paragraph_text(p).replace("\xa0", " ").strip()
    if text:
        runs = [r for r in p.runs if r.text.strip()]
        letters = [c for c in text if c.isalpha()]
        out.append(Block(
            text=text.replace("\t", " ⇥ ").replace("\n", " ⏎ "),
            where=where,
            style=p.style.name if p.style is not None else "",
            bullet=_is_bullet(p),
            bold=bool(runs) and all(_run_attr(r, p, "bold") for r in runs),
            italic=bool(runs) and all(_run_attr(r, p, "italic") for r in runs),
            caps=(len(letters) > 1 and all(c.isupper() for c in letters))
                 or any(r.font.all_caps for r in runs),
            size=max((s for s in (_size(r, p) for r in runs) if s), default=None),
            align=_alignment(p),
            rule=_has_bottom_rule(p),
            links=[(h.text, h.url) for h in p.hyperlinks if h.url],
            para=p,
        ))
    # text boxes anchored in this paragraph (skip the duplicate fallback copy)
    for tb in p._p.iter(qn("w:txbxContent")):
        if any(a.tag == MC_FALLBACK for a in tb.iterancestors()):
            continue
        _walk(tb, p, "textbox", out)


def _walk(container_el, parent, where: str, out: list[Block]) -> None:
    for item in _iter_items(container_el, parent):
        if isinstance(item, Paragraph):
            _add_paragraph(item, where, out)
        else:
            for ri, row in enumerate(item.rows):
                seen = set()
                for ci, cell in enumerate(row.cells):
                    if cell._tc in seen:               # skip merged duplicates
                        continue
                    seen.add(cell._tc)
                    _walk(cell._tc, cell, f"table r{ri}c{ci}", out)


def extract_blocks(doc) -> list[Block]:
    out: list[Block] = []
    for i, section in enumerate(doc.sections):          # headers first: contact info lives here
        if i == 0 or not section.header.is_linked_to_previous:
            _walk(section.header._element, section.header, "header", out)
    _walk(doc.element.body, doc, "body", out)
    for i, section in enumerate(doc.sections):
        if i == 0 or not section.footer.is_linked_to_previous:
            _walk(section.footer._element, section.footer, "footer", out)
    return out


def render(blocks: list[Block]) -> str:
    sizes = [b.size for b in blocks if b.size]
    body = median(sizes) if sizes else None
    lines = []
    for i, b in enumerate(blocks):
        flags = []
        if b.where != "body":
            flags.append(b.where)
        if b.bullet:
            flags.append("BULLET")
        if b.bold:
            flags.append("BOLD")
        if b.italic:
            flags.append("ITALIC")
        if b.caps:
            flags.append("CAPS")
        if b.rule:
            flags.append("RULE_BELOW")
        if b.align in ("center", "right"):
            flags.append(b.align.upper())
        if body and b.size and b.size >= body + 1.5:
            flags.append(f"LARGE({b.size:g}pt)")
        if b.style and not b.style.lower().startswith(("normal", "list paragraph")):
            flags.append(f"style={b.style}")
        line = f"[{i}] {' '.join(flags)} | {b.text}"
        if b.links:
            line += " | LINKS: " + "; ".join(f"{t} -> {u}" for t, u in b.links)
        lines.append(line)
    return "\n".join(lines)


def dump(path: str) -> str:
    blocks = extract_blocks(Document(path))
    return render(blocks) if blocks else ""


if __name__ == "__main__":
    print(dump(sys.argv[1]))

